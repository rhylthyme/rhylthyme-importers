"""
Slide deck importer - converts PowerPoint (.pptx) files to Rhylthyme presentation timelines.
"""

import re
from pathlib import Path
from typing import Dict, Any, List, Optional

from pptx import Presentation
from pptx.util import Inches

from .base import BaseImporter, ImportResult


# Speaking rate assumptions
WORDS_PER_MINUTE = 150
BASE_SECONDS_PER_SLIDE = 120
MAX_SECONDS_PER_SLIDE = 600


def _extract_text_from_slide(slide) -> tuple[str, str, str]:
    """Extract title, body text, and speaker notes from a slide.

    Returns (title, body_text, notes).
    """
    title = ""
    body_parts = []

    if slide.shapes.title is not None:
        title = slide.shapes.title.text.strip()

    for shape in slide.shapes:
        if shape.has_text_frame:
            # Skip the title shape (already captured)
            if slide.shapes.title is not None and shape == slide.shapes.title:
                continue
            text = shape.text_frame.text.strip()
            if text:
                body_parts.append(text)

    body_text = "\n".join(body_parts)

    # If no title shape but body is very short, promote it to title
    if not title and body_text:
        body_words = len(body_text.split())
        if body_words <= 5:
            title = body_text
            body_text = ""

    notes = ""
    try:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
    except Exception:
        pass

    return title, body_text, notes


def _parse_duration_hint(notes: str) -> Optional[int]:
    """Parse timing hints from speaker notes like '2 min', '5 minutes', '90 seconds'."""
    if not notes:
        return None

    patterns = [
        r'(\d+)\s*min(?:ute)?s?\b',
        r'(\d+)\s*sec(?:ond)?s?\b',
    ]

    for pattern in patterns:
        match = re.search(pattern, notes, re.IGNORECASE)
        if match:
            value = int(match.group(1))
            if 'sec' in pattern:
                return min(value, MAX_SECONDS_PER_SLIDE)
            else:
                return min(value * 60, MAX_SECONDS_PER_SLIDE)

    return None


def _estimate_duration(body_text: str, notes: str) -> tuple[int, int, int]:
    """Estimate slide duration as (min, default, max) in seconds.

    Priority: explicit timing hint in notes > word-count estimate.
    Returns variable bounds: min ~60% of default, max ~175% of default.
    """
    hint = _parse_duration_hint(notes)
    if hint is not None:
        default = hint
    else:
        word_count = len(body_text.split()) if body_text else 0
        # ~1 second per word on top of base time
        default = BASE_SECONDS_PER_SLIDE + word_count

    default = min(default, MAX_SECONDS_PER_SLIDE)
    # You could rush through a slide in ~60% of the default time
    min_secs = max(30, int(default * 0.6))
    # You might linger, take questions, or elaborate — up to ~175%
    max_secs = min(MAX_SECONDS_PER_SLIDE, int(default * 1.75))

    return min_secs, default, max_secs


def _is_section_header(title: str, body_text: str) -> bool:
    """Detect section header slides (title-only, no substantial body)."""
    if not title:
        return False
    # Section header = has a title but no meaningful body text
    body_words = len(body_text.split()) if body_text else 0
    return body_words <= 3


class SlideDeckImporter(BaseImporter):
    """Import PowerPoint slide decks (.pptx) as presentation timelines."""

    name = "slidedeck"
    description = "Import slide decks (.pptx) as presentation timelines"
    supported_domains: List[str] = []  # File-based, not URL-based

    def can_import(self, url_or_query: str) -> bool:
        return url_or_query.lower().endswith('.pptx')

    def import_from_file(self, file_path: str) -> ImportResult:
        """Import a .pptx file and convert to a Rhylthyme program."""
        try:
            prs = Presentation(file_path)
        except Exception as e:
            return ImportResult(
                success=False,
                error=f"Failed to open PowerPoint file: {e}",
                source_type="slidedeck",
            )

        filename = Path(file_path).stem

        # Extract core properties for the presentation title
        pptx_title = None
        try:
            if prs.core_properties and prs.core_properties.title:
                pptx_title = prs.core_properties.title.strip()
        except Exception:
            pass
        presentation_name = pptx_title or filename

        # Extract all slides
        slides_data = []
        for i, slide in enumerate(prs.slides):
            title, body_text, notes = _extract_text_from_slide(slide)
            # Skip blank slides
            if not title and not body_text:
                continue
            slides_data.append({
                "index": i + 1,
                "title": title,
                "body_text": body_text,
                "notes": notes,
                "is_section_header": _is_section_header(title, body_text),
            })

        if not slides_data:
            return ImportResult(
                success=False,
                error="No slides with content found in the PowerPoint file.",
                source_type="slidedeck",
            )

        # Decide track structure: multi-track if section headers exist
        has_sections = any(s["is_section_header"] for s in slides_data)
        tracks = self._build_tracks(slides_data, has_sections)

        program = {
            "schemaVersion": "0.1.0",
            "programId": self.generate_program_id(f"presentation-{filename}"),
            "name": presentation_name,
            "description": f"Imported from {Path(file_path).name}",
            "environmentType": "event",
            "metadata": {
                "programType": "presentation",
            },
            "tracks": tracks,
            "resourceConstraints": [
                {"task": "presentation", "maxConcurrent": 1}
            ],
        }

        return ImportResult(
            success=True,
            program=program,
            source_type="slidedeck",
        )

    def _build_tracks(
        self, slides_data: List[Dict], has_sections: bool
    ) -> List[Dict[str, Any]]:
        """Build track(s) from slide data.

        Tracks are sequential — each track's first step chains after the
        last step of the previous track so the presentation flows linearly.
        """
        if not has_sections:
            track = self._build_single_track("presentation", "Presentation", slides_data)
            # Set step-level priority based on position:
            # first 25% = 50 (important), middle 50% = 100 (normal), last 25% = 150 (nice to have)
            n = len(track["steps"])
            for i, step in enumerate(track["steps"]):
                frac = i / n if n > 0 else 0
                if frac < 0.25:
                    step["priority"] = 50
                elif frac < 0.75:
                    step["priority"] = 100
                else:
                    step["priority"] = 150
            return [track]

        # Multi-track: split on section header slides
        sections: List[tuple[str, List[Dict]]] = []
        current_section_name = "Introduction"
        current_slides: List[Dict] = []
        section_index = 0

        for slide in slides_data:
            if slide["is_section_header"]:
                # Flush previous section
                if current_slides:
                    sections.append((current_section_name, current_slides))
                    section_index += 1
                    current_slides = []
                current_section_name = slide["title"] or f"Section {section_index + 1}"
            else:
                current_slides.append(slide)

        # Flush last section
        if current_slides:
            sections.append((current_section_name, current_slides))

        # If no content slides ended up in sections (all were section headers), make one track
        if not sections:
            return [self._build_single_track("presentation", "Presentation", slides_data)]

        # Build tracks, chaining each track's first step after the previous track's last step
        tracks = []
        prev_last_step_id: Optional[str] = None
        for idx, (section_name, slides) in enumerate(sections):
            track_id = self.generate_program_id(section_name)
            track = self._build_single_track(
                track_id, section_name, slides, chain_after=prev_last_step_id
            )
            # Set track priority based on position (first = highest priority)
            track["priority"] = idx + 1
            tracks.append(track)
            prev_last_step_id = track["steps"][-1]["stepId"]

        return tracks

    def _build_single_track(
        self,
        track_id: str,
        track_name: str,
        slides: List[Dict],
        chain_after: Optional[str] = None,
    ) -> Dict[str, Any]:
        """Build a single track from a list of slide data dicts.

        Args:
            chain_after: stepId from a previous track. If set, this track's
                         first step uses afterStep to chain sequentially.
        """
        steps = []
        for i, slide in enumerate(slides):
            title = slide["title"] or f"Slide {slide['index']}"
            step_id = self.generate_program_id(f"slide-{slide['index']}-{title}")
            min_secs, default_secs, max_secs = _estimate_duration(
                slide["body_text"], slide["notes"]
            )

            # Build description from body text and notes
            desc_parts = []
            if slide["body_text"]:
                desc_parts.append(slide["body_text"])
            if slide["notes"]:
                desc_parts.append(f"Speaker notes: {slide['notes']}")
            description = "\n\n".join(desc_parts) if desc_parts else None

            step: Dict[str, Any] = {
                "stepId": step_id,
                "name": title,
                "task": "presentation",
                "duration": {
                    "type": "variable",
                    "minSeconds": min_secs,
                    "defaultSeconds": default_secs,
                    "maxSeconds": max_secs,
                },
            }

            if i == 0:
                if chain_after:
                    step["startTrigger"] = {"type": "afterStep", "stepId": chain_after}
                else:
                    step["startTrigger"] = {"type": "programStart"}
            else:
                prev_step_id = steps[i - 1]["stepId"]
                step["startTrigger"] = {"type": "afterStep", "stepId": prev_step_id}

            if description:
                step["description"] = description

            steps.append(step)

        return {
            "trackId": track_id,
            "name": track_name,
            "steps": steps,
        }

    def import_from_url(self, url: str) -> ImportResult:
        return self.import_from_file(url)

    def search(self, query: str) -> List[Dict[str, Any]]:
        return []
